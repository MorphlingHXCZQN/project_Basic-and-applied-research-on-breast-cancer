"""Create project proposal documents."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from .gpt_client import GPTClient


class ProposalWriter:
    """Generate project proposals and export them to Word."""

    def __init__(self, gpt: GPTClient) -> None:
        self._gpt = gpt

    def draft_proposal(self, outline: str, word_count: int) -> str:
        return self._gpt.draft_proposal(outline, word_count)

    def save_to_word(self, content: str, output_path: Path) -> Path:
        document = Document()
        for paragraph in content.splitlines():
            if paragraph.strip():
                document.add_paragraph(paragraph)
            else:
                document.add_paragraph("")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
        return output_path
